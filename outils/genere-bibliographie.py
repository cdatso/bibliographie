# -*- coding: ascii -*-
"""Genere index.html (bibliographie de Christo Datso) depuis le classeur de reference.

Le classeur est lu en LECTURE SEULE (jamais ouvert en ecriture). Une seule ecriture : index.html du depot.
Usage :
  python genere-bibliographie.py --date "AAAA-MM-JJ HH:MM"      genere index.html
  python genere-bibliographie.py --simuler                       affiche les comptes, n'ecrit rien
  python genere-bibliographie.py --verifier                      regenere en memoire, compare a index.html
Options : --classeur <chemin du .xlsx>   --depot <dossier du depot (defaut : parent de outils)>
Codes de sortie : 0 ok ; 1 --verifier : page differente ; 2 donnees refusees ou usage incorrect.
Ce fichier est en ASCII pur : les lettres accentuees y sont ecrites {e} (= e aigu) ou en entites HTML.
"""
import sys, os, re, glob, html, hashlib, argparse, collections
import openpyxl

SHEET = "sources_bibliographie_christo_d"


def fx(s):
    return s.replace("{e}", chr(233))


E_OK = [fx(s) for s in ("nouvelle", "chronique / recension", "novella", "recueil de nouvelles", "essai", "article",
                        "article (autre auteur)", "po{e}sie", "lettre / article", "aide de jeu",
                        "fichier de fiches", "roman")]
F_OK = [fx(s) for s in ("papier + online", "papier", "online", "epub", "in{e}dit (usage priv{e} JdR)",
                        "audio online", "in{e}dit (manuscrit)")]
INEDIT = fx("in{e}dit")
AUTRE_AUTEUR = "article (autre auteur)"

# Rubriques du sommaire : (id d'ancre, libelle HTML, test sur (E, F)). Un texte peut figurer dans deux rubriques.
GROUPES = [
    ("fiction", "Fiction", lambda e, f: e in ("nouvelle", "novella", "recueil de nouvelles", "roman")),
    ("chroniques", "Chroniques et recensions", lambda e, f: e == "chronique / recension"),
    ("essais", "Essais et articles", lambda e, f: e in ("essai", "article", "lettre / article")),
    ("poesie", "Po&eacute;sie", lambda e, f: e == fx("po{e}sie")),
    ("audio", "Audio", lambda e, f: f == "audio online"),
]
TYPES_PUBLIES = set(["nouvelle", "novella", "recueil de nouvelles", "roman", "chronique / recension", "essai",
                     "article", "lettre / article", fx("po{e}sie")])

CSS = """
    :root {
      color-scheme: light dark;
      --papier: #faf8f4;
      --encre: #1a1a1a;
      --discret: #555;
      --or: #b8993f;
      --or-texte: #7a6120;
    }
    @media (prefers-color-scheme: dark) {
      :root {
        --papier: #14140f;
        --encre: #f0ede6;
        --discret: #b5b0a6;
        --or-texte: #d4b862;
      }
    }
    * { box-sizing: border-box; margin: 0; padding: 0; }
    body {
      min-height: 100vh;
      font-family: system-ui, sans-serif;
      background-color: var(--papier);
      color: var(--encre);
      padding: 2.5rem 1.25rem 2rem;
      line-height: 1.5;
    }
    main { max-width: 44rem; margin: 0 auto; }
    header { text-align: center; margin-bottom: 2rem; }
    h1, h2 { font-family: Georgia, "Times New Roman", serif; font-weight: 400; }
    h1 { font-size: clamp(2.2rem, 7vw, 3.4rem); line-height: 1.15; margin-bottom: 1rem; }
    h2 { font-size: 1.5rem; margin: 2rem 0 0.6rem; padding-bottom: 0.2rem; border-bottom: 1px solid var(--or); }
    .presentation {
      font-family: Georgia, "Times New Roman", serif;
      font-size: 1.1rem;
      line-height: 1.6;
      color: var(--discret);
      max-width: 34rem;
      margin: 0 auto;
    }
    .retour { display: inline-block; margin-bottom: 1.5rem; font-size: 0.88rem; color: var(--discret); }
    a { color: var(--or-texte); }
    a:hover { color: var(--encre); }
    a:focus-visible { outline: 3px solid var(--or-texte); outline-offset: 3px; }
    ul { list-style: none; }
    nav.sommaire { margin-bottom: 1.5rem; text-align: center; }
    nav.sommaire ul { display: flex; flex-wrap: wrap; justify-content: center; gap: 0.5rem 1.2rem; font-size: 0.95rem; }
    nav.annees { font-size: 0.85rem; color: var(--discret); text-align: center; margin-bottom: 0.5rem; }
    nav.annees a { margin: 0 0.25rem; white-space: nowrap; }
    section.liste li { padding: 0.45rem 0; border-bottom: 1px dotted var(--discret); font-size: 0.95rem; }
    section.liste li:target { background: rgba(184, 153, 63, 0.18); }
    section.liste cite { font-family: Georgia, "Times New Roman", serif; font-style: italic; }
    .meta { color: var(--discret); }
    .par-type li { font-size: 0.9rem; padding: 0.15rem 0; }
    .par-type h3 { font-family: Georgia, "Times New Roman", serif; font-weight: 400; font-size: 1.15rem; margin: 1.2rem 0 0.4rem; }
    .note { font-size: 0.85rem; color: var(--discret); }
    footer { margin-top: 3rem; text-align: center; font-size: 0.8rem; color: var(--discret); }
    footer p { margin-bottom: 0.3rem; }
    footer a { color: inherit; }
"""


class Refus(Exception):
    pass


def cellule(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return " ".join(str(v).split())


def lire_classeur(chemin):
    """Lecture seule : octets pour l'empreinte, puis openpyxl en read_only. Retourne (empreinte, lignes)."""
    with open(chemin, "rb") as fh:
        empreinte = hashlib.sha256(fh.read()).hexdigest()
    wb = openpyxl.load_workbook(chemin, read_only=True, data_only=True)
    try:
        if SHEET not in wb.sheetnames:
            raise Refus("feuille absente : %s" % SHEET)
        ws = wb[SHEET]
        brut = list(ws.iter_rows(min_row=2, max_col=10, values_only=True))
    finally:
        wb.close()
    lignes = []
    ids = set()
    for r in brut:
        if not any(v is not None and str(v).strip() for v in r):
            continue
        id_ = r[0]
        if isinstance(id_, float) and id_.is_integer():
            id_ = int(id_)
        if not isinstance(id_, int) or isinstance(id_, bool):
            raise Refus("ID non entier : %r" % (id_,))
        if id_ in ids:
            raise Refus("ID en double : %d" % id_)
        ids.add(id_)
        e, f, h = cellule(r[4]), cellule(r[5]), cellule(r[7])
        if e not in E_OK:
            raise Refus("ID %d : type de texte hors liste : %r" % (id_, e))
        if f not in F_OK:
            raise Refus("ID %d : type de publication hors liste : %r" % (id_, f))
        if h and not h.startswith("http"):
            raise Refus("ID %d : URL ne commence pas par http : %r" % (id_, h))
        g = r[6]
        if isinstance(g, float) and g.is_integer():
            g = int(g)
        if g is not None and cellule(g) == "":
            g = None
        if g is not None and (not isinstance(g, int) or isinstance(g, bool) or not 1900 <= g <= 2100):
            raise Refus("ID %d : annee invalide : %r" % (id_, g))
        d = cellule(r[3])
        if not d:
            raise Refus("ID %d : titre du texte vide" % id_)
        lignes.append({"id": id_, "b": cellule(r[1]), "c": cellule(r[2]), "d": d, "e": e, "f": f, "g": g, "h": h})
    return empreinte, lignes


def perimetre(lignes):
    pub, exclus = [], []
    for l in lignes:
        if l["f"].startswith(INEDIT) or l["e"] == AUTRE_AUTEUR:
            exclus.append(l)
        else:
            if l["e"] not in TYPES_PUBLIES:
                raise Refus("ID %d : type de texte publie sans rubrique : %r" % (l["id"], l["e"]))
            pub.append(l)
    pub.sort(key=lambda l: (l["g"] is None, -(l["g"] or 0), l["id"]))
    return pub, exclus


def esc(s):
    return html.escape(s, quote=True)


def entree(l):
    titre = esc(l["d"])
    if l["h"]:
        titre = '<a href="%s" rel="noopener">%s</a>' % (esc(l["h"]), titre)
    s = "<cite>%s</cite> &mdash; %s" % (titre, esc(l["e"]))
    src = [x for x in (l["b"], l["c"]) if x]
    if src:
        s += ' <span class="meta">&middot; dans %s</span>' % esc(", ".join(src))
    s += ' <span class="meta">(%s)</span>' % esc(l["f"])
    return '<li id="e-%d">%s</li>' % (l["id"], s)


def construire(empreinte, pub, date):
    annees = collections.OrderedDict()
    for l in pub:
        annees.setdefault(l["g"], []).append(l)
    vues = [l["g"] for l in pub if l["g"] is not None]
    plage = "%d-%d" % (min(vues), max(vues)) if vues else ""
    groupes = [(gid, lib, [l for l in pub if test(l["e"], l["f"])]) for gid, lib, test in GROUPES]
    groupes = [g for g in groupes if g[2]]
    o = []
    o.append("<!doctype html>")
    o.append('<html lang="fr">')
    o.append("<head>")
    o.append('  <meta charset="utf-8">')
    o.append('  <meta name="viewport" content="width=device-width, initial-scale=1">')
    o.append("  <title>Bibliographie &mdash; Christo Datso</title>")
    o.append('  <meta name="description" content="Bibliographie de Christo Datso : nouvelles, novellas, recueils, '
             'chroniques, essais et po&eacute;sie publi&eacute;s, avec liens vers les pages des &eacute;diteurs.">')
    o.append("  <style>" + CSS + "  </style>")
    o.append("</head>")
    o.append("<body>")
    o.append("  <main>")
    o.append("    <header>")
    o.append('      <a class="retour" href="/">&larr; Christo Datso</a>')
    o.append("      <h1>Bibliographie</h1>")
    o.append('      <p class="presentation">Textes publi&eacute;s de Christo Datso, %s&nbsp;: nouvelles, novellas, '
             "recueils, chroniques, essais.</p>" % plage)
    o.append("    </header>")
    o.append('    <nav class="sommaire" aria-label="Sommaire par type">')
    o.append("      <ul>")
    for gid, lib, ls in groupes:
        o.append('        <li><a href="#t-%s">%s (%d)</a></li>' % (gid, lib, len(ls)))
    o.append("      </ul>")
    o.append("    </nav>")
    o.append('    <nav class="annees" aria-label="Sommaire par ann&eacute;e">')
    liens = []
    for a in annees:
        if a is None:
            liens.append('<a href="#a-sans-date">Sans date</a>')
        else:
            liens.append('<a href="#a-%d">%d</a>' % (a, a))
    o.append("      " + " ".join(liens))
    o.append("    </nav>")
    o.append('    <section class="liste" aria-label="Liste par ann&eacute;e">')
    for a, ls in annees.items():
        if a is None:
            o.append('      <h2 id="a-sans-date">Sans date</h2>')
        else:
            o.append('      <h2 id="a-%d">%d</h2>' % (a, a))
        o.append("      <ul>")
        for l in ls:
            o.append("        " + entree(l))
        o.append("      </ul>")
    o.append("    </section>")
    o.append('    <section class="par-type" aria-label="Index par type">')
    o.append("      <h2>Index par type</h2>")
    o.append('      <p class="note">Un m&ecirc;me texte peut figurer dans deux rubriques. '
             "Chaque titre renvoie &agrave; son entr&eacute;e dans la liste ci-dessus.</p>")
    for gid, lib, ls in groupes:
        o.append('      <h3 id="t-%s">%s (%d)</h3>' % (gid, lib, len(ls)))
        o.append("      <ul>")
        for l in ls:
            o.append('        <li><a href="#e-%d">%s</a></li>' % (l["id"], esc(l["d"])))
        o.append("      </ul>")
    o.append("    </section>")
    o.append("  </main>")
    o.append("  <footer>")
    o.append('    <p>Christo Datso &middot; <a href="mailto:analyses@cdatso.be">analyses@cdatso.be</a></p>')
    o.append("    <p>g&eacute;n&eacute;r&eacute; le %s depuis le classeur de r&eacute;f&eacute;rence (SHA-256 %s)</p>"
             % (esc(date), empreinte[:8]))
    o.append("    <p>%d textes publi&eacute;s</p>" % len(pub))
    o.append("  </footer>")
    o.append("</body>")
    o.append("</html>")
    return "\n".join(o) + "\n"


def classeur_par_defaut():
    base = glob.glob(r"D:\documentation\Base de Donn*es auteur Christo Datso")
    if not base:
        raise Refus("dossier du classeur introuvable sur D:")
    trouves = glob.glob(os.path.join(base[0], "sources_bibliographie_christo_datso (mises * jours).xlsx"))
    if not trouves:
        raise Refus("classeur introuvable")
    return trouves[0]


def main():
    sys.stdout.reconfigure(encoding="utf-8")
    ap = argparse.ArgumentParser(description="Genere index.html depuis le classeur (lecture seule).")
    ap.add_argument("--classeur")
    ap.add_argument("--depot", default=os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
    ap.add_argument("--date", help='AAAA-MM-JJ HH:MM (lue d\'un Get-Date, jamais estimee)')
    ap.add_argument("--verifier", action="store_true")
    ap.add_argument("--simuler", action="store_true")
    a = ap.parse_args()
    try:
        classeur = a.classeur or classeur_par_defaut()
        empreinte, lignes = lire_classeur(classeur)
        pub, exclus = perimetre(lignes)
        cible = os.path.join(a.depot, "index.html")
        if a.simuler:
            print("classeur : %s" % classeur)
            print("SHA-256 : %s" % empreinte)
            print("lues : %d | publiees : %d | exclues : %d (ID %s)" % (
                len(lignes), len(pub), len(exclus), ", ".join(str(l["id"]) for l in exclus)))
            for gid, lib, test in GROUPES:
                print("rubrique %s : %d" % (gid, sum(1 for l in pub if test(l["e"], l["f"]))))
            print("sans annee : %d | avec URL : %d" % (sum(1 for l in pub if l["g"] is None),
                                                       sum(1 for l in pub if l["h"])))
            return 0
        if a.verifier:
            if not os.path.isfile(cible):
                print("index.html absent : %s" % cible)
                return 1
            actuel = open(cible, "r", encoding="utf-8", newline="").read()
            date = a.date
            if not date:
                m = re.search(r"g&eacute;n&eacute;r&eacute; le (.+?) depuis", actuel)
                if not m:
                    print("date introuvable dans index.html")
                    return 1
                date = html.unescape(m.group(1))
            if construire(empreinte, pub, date) == actuel:
                print("verifier : index.html egal (%d textes)" % len(pub))
                return 0
            print("verifier : index.html DIFFERENT de la generation")
            return 1
        if not a.date:
            raise Refus("--date requis")
        page = construire(empreinte, pub, a.date)
        with open(cible, "w", encoding="utf-8", newline="\n") as fh:
            fh.write(page)
        print("ecrit : %s | %d textes | %d octets" % (cible, len(pub), len(page.encode("utf-8"))))
        return 0
    except Refus as ex:
        print("REFUS : %s" % ex)
        return 2


if __name__ == "__main__":
    sys.exit(main())
